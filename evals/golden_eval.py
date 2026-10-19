#!/usr/bin/env python3
"""Golden-50 evaluation (no model calls).

  python evals/golden_eval.py convert --xlsx <labeled.xlsx>     # human labels -> evals/golden_50_labels.csv
  python evals/golden_eval.py evaluate --records runs/golden/records.jsonl

Human labels are only ever read here, after predictions exist; they are never
model inputs, prompt examples, thresholds or issue-discovery inputs.
"""
import argparse
import csv
import json
import sys
from collections import Counter, defaultdict
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))
from pipeline.io_utils import read_jsonl, write_json  # noqa: E402

LABELS = ROOT / "evals" / "golden_50_labels.csv"
SOURCE = ROOT / "data" / "golden_50_to_label.csv"
COLS = ["review_id", "topic", "intent", "severity", "sentiment", "entities", "evidence_quote", "needs_review", "ambiguous", "notes"]
SENT_TOL = 0.5  # predeclared sentiment tolerance


def convert(xlsx):
    import openpyxl
    src = {r["review_id"]: r for r in csv.DictReader(open(SOURCE, newline="", encoding="utf-8"))}
    ws = openpyxl.load_workbook(xlsx, read_only=True)["Label Here"]
    hdr = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    out = []
    for vals in ws.iter_rows(min_row=2, values_only=True):
        r = dict(zip(hdr, vals))
        if not r.get("review_id"):
            continue
        text = src[r["review_id"]]["review_text"]  # always the original text, never the sheet copy
        q = r.get("evidence_quote") or text
        assert q in text, f"quote not in source for {r['review_id']}"
        ents = [e.strip() for e in str(r.get("entities") or "").split(",") if e.strip()]
        out.append({"review_id": r["review_id"], "topic": r["topic"], "intent": r["intent"], "severity": int(r["severity"]),
                    "sentiment": float(r["sentiment"]), "entities": ";".join(ents), "evidence_quote": q,
                    "needs_review": str(r.get("needs_review")).upper() == "TRUE", "ambiguous": str(r.get("ambiguous")).lower() == "yes",
                    "notes": r.get("my_notes") or ""})
    assert len(out) == 50 and set(x["review_id"] for x in out) == set(src)
    with open(LABELS, "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=COLS)
        w.writeheader()
        w.writerows(out)
    print(f"wrote {LABELS} ({len(out)} human-labeled rows)")


def evaluate(records_path, out_dir):
    human = {r["review_id"]: r for r in csv.DictReader(open(LABELS, newline="", encoding="utf-8"))}
    texts = {r["review_id"]: r["review_text"] for r in csv.DictReader(open(SOURCE, newline="", encoding="utf-8"))}
    preds = {r["review_id"]: r for r in read_jsonl(records_path) if r["review_id"] in human}
    cases, conf = [], {f: Counter() for f in ("topic", "intent")}
    agg = Counter()
    per_topic = defaultdict(Counter)
    for rid, h in human.items():
        p = preds.get(rid)
        ok = p is not None and p["status"] == "completed"
        hs, hsev = float(h["sentiment"]), int(h["severity"])
        row = {"review_id": rid, "text": texts[rid][:200], "ambiguous": h["ambiguous"], "human_notes": h["notes"],
               "h_topic": h["topic"], "p_topic": p.get("topic") if ok else "<missing>",
               "h_intent": h["intent"], "p_intent": p.get("intent") if ok else "<missing>",
               "h_severity": hsev, "p_severity": p.get("severity") if ok else "",
               "h_sentiment": hs, "p_sentiment": p.get("sentiment") if ok else "",
               "h_needs_review": h["needs_review"], "p_needs_review": p.get("needs_review") if ok else ""}
        row["topic_ok"] = ok and p["topic"] == h["topic"]
        row["intent_ok"] = ok and p["intent"] == h["intent"]
        row["severity_ok"] = ok and p["severity"] == hsev
        row["sentiment_ok"] = ok and abs(p["sentiment"] - hs) <= SENT_TOL
        if ok:
            row["quote_exact_substring"] = p["evidence_quote"] in texts[rid]
            hents = {e.lower() for e in h["entities"].split(";") if e}
            row["entities_pred"] = ";".join(p["entities"])
            row["entities_unsupported"] = ";".join(e for e in p["entities"] if e.lower() not in texts[rid].lower())
            row["entities_not_in_human"] = ";".join(e for e in p["entities"] if e.lower() not in hents)
            agg["sev_abs_err"] += abs(p["severity"] - hsev)
            agg["sent_abs_err"] += abs(p["sentiment"] - hs)
            agg["valid"] += 1
            agg["quote_exact"] += row["quote_exact_substring"]
            agg["unsupported_entities"] += bool(row["entities_unsupported"])
        for f in ("topic", "intent", "severity", "sentiment"):
            agg[f + "_ok"] += row[f + "_ok"]
        conf["topic"][(h["topic"], row["p_topic"])] += 1
        conf["intent"][(h["intent"], row["p_intent"])] += 1
        per_topic[h["topic"]]["human"] += 1
        per_topic[row["p_topic"]]["predicted"] += 1
        per_topic[h["topic"]]["correct"] += row["topic_ok"]
        agg["ambiguous"] += h["ambiguous"] == "True"
        if h["ambiguous"] == "True":
            agg["amb_topic_ok"] += row["topic_ok"]
        # needs_review as a prediction of human ambiguity / needs_review
        hflag = h["ambiguous"] == "True" or h["needs_review"] == "True"
        pflag = bool(ok and p["needs_review"])
        agg["nr_tp"] += hflag and pflag
        agg["nr_fp"] += pflag and not hflag
        agg["nr_fn"] += hflag and not pflag
        cases.append(row)
    n = len(human)
    summary = {"n": n, "valid_predictions": agg["valid"],
               "topic_agreement_pct": f"{100 * agg['topic_ok'] / n:.1f}", "intent_agreement_pct": f"{100 * agg['intent_ok'] / n:.1f}",
               "severity_exact_pct": f"{100 * agg['severity_ok'] / n:.1f}",
               "severity_mae": f"{agg['sev_abs_err'] / max(1, agg['valid']):.2f}",
               "sentiment_within_tolerance_pct": f"{100 * agg['sentiment_ok'] / n:.1f}", "sentiment_tolerance": SENT_TOL,
               "sentiment_mae": f"{agg['sent_abs_err'] / max(1, agg['valid']):.2f}",
               "ambiguous_cases": agg["ambiguous"], "topic_agreement_on_ambiguous": f"{agg['amb_topic_ok']}/{agg['ambiguous']}",
               "quote_exact_substring": f"{agg['quote_exact']}/{agg['valid']}",
               "records_with_unsupported_entities": agg["unsupported_entities"],
               "needs_review_vs_human_flag": {"tp": agg["nr_tp"], "fp": agg["nr_fp"], "fn": agg["nr_fn"]},
               "per_topic": {k: dict(v) for k, v in sorted(per_topic.items())}}
    confusion = {f: [{"human": h, "predicted": p, "count": c} for (h, p), c in sorted(v.items())] for f, v in conf.items()}
    out_dir.mkdir(parents=True, exist_ok=True)
    write_json(out_dir / "golden_results.json", {"summary": summary, "confusion": confusion, "records": str(records_path),
                                                 "note": "50 cases is a small diagnostic sample, not a population accuracy estimate."})
    with open(out_dir / "golden_cases.csv", "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=list(cases[0].keys()))
        w.writeheader()
        w.writerows(cases)
    print(json.dumps(summary, indent=1))


if __name__ == "__main__":
    p = argparse.ArgumentParser()
    sub = p.add_subparsers(dest="cmd", required=True)
    c = sub.add_parser("convert"); c.add_argument("--xlsx", type=Path, required=True)
    e = sub.add_parser("evaluate"); e.add_argument("--records", type=Path, required=True)
    e.add_argument("--out-dir", type=Path, default=ROOT / "evals")
    a = p.parse_args()
    convert(a.xlsx) if a.cmd == "convert" else evaluate(a.records, a.out_dir)
